#!/usr/bin/env python3
"""
depseudonymize.py — deterministische Rueckumwandlung eines PSEUDO-DOCX in Klartext.

Skill-Version 1.1, unterstuetzt Legendenschema v1 (Grundform) und v2 (Positionsindex).

Aufruf:
    python depseudonymize.py \
        --input Dokument_PSEUDO.docx \
        --legend MANDAT-XY_LEGENDE_v1.json \
        --outdir ./output

Erzeugt:
    <name>_KLAR.docx   — DOCX mit rueckuebersetzten Codes

Regeln
------
- Jeder Code der Form [KATEGORIE_NN] wird ersetzt:
  * Schema v2: durch die stellenbezogene 'originalform' aus der Legende
    (positionsgetreu, Cursor je Code); wortgleicher Roundtrip moeglich.
  * Schema v1: durch die 'grundform' aus der Legende; Bericht enthaelt
    dann eine Warnung, dass wortgleicher Roundtrip nicht garantiert ist.
- Codes ohne Legendeneintrag bleiben stehen und werden im Bericht als
  Fehler gemeldet.
- Die Kennzeichnungszeile aus Phase 2 wird entfernt.
- Keine Umformulierungen, keine Zusaetze.
- Iterationsreihenfolge des Dokuments (Absaetze, dann Tabellen; danach
  Kopf-/Fusszeilen aller Sections) MUSS dieselbe sein wie in
  pseudonymize.py, sonst laeuft der Cursor auseinander.
"""

from __future__ import annotations
import argparse
import json
import re
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    sys.stderr.write(
        "FEHLER: python-docx nicht installiert. "
        "Bitte 'pip install -r requirements.txt' ausfuehren.\n"
    )
    sys.exit(2)


CODE_PATTERN = re.compile(
    r"\[(PERSON|FIRMA|BEHOERDE|ORT|ANSCHRIFT|KONTAKT|STNR|REGISTER|"
    r"KONTO|AZ|GEBDAT|OBJEKT|SONSTIGES)_\d{2,3}\]"
)

MARKER_PATTERN = re.compile(
    r"^\[Pseudonymisiertes Dokument · Alias .+? · "
    r"Legende v\d+ · Codes nicht aufl(ö|oe)sen\]$"
)


def load_legend(path: Path) -> tuple[int, dict[str, dict]]:
    """
    Laedt die Legende und gibt (schema_version, eintrag_by_code) zurueck.
    Fuer Schema v1 werden die Eintraege intern auf v2-aehnliche Form gebracht
    (Feld 'vorkommen' bleibt dann leer, Fallback auf Grundform).
    """
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    schema_version = int(data.get("version", 1))
    eintrag_by_code: dict[str, dict] = {}
    for eintrag in data.get("eintraege", []):
        code = eintrag["code"]
        e = {
            "grundform": eintrag["grundform"],
            "vorkommen": eintrag.get("vorkommen", []) if schema_version >= 2 else [],
            "kategorie": eintrag.get("kategorie", ""),
        }
        eintrag_by_code[code] = e
    return schema_version, eintrag_by_code


def make_reverse_replacer(schema_version: int, eintrag_by_code: dict[str, dict],
                          cursor: dict[str, int],
                          counter: dict[str, int],
                          missing: set[str],
                          fallback_count: dict[str, int]):
    """Baut die re-sub-Callback-Funktion mit gebundenen Statuscontainern."""
    def repl(match: re.Match) -> str:
        code_full = match.group(0)     # z. B. "[PERSON_01]"
        code_key = code_full[1:-1]     # "PERSON_01"
        if code_key not in eintrag_by_code:
            missing.add(code_full)
            return code_full
        eintrag = eintrag_by_code[code_key]
        counter[code_full] = counter.get(code_full, 0) + 1

        if schema_version >= 2 and eintrag["vorkommen"]:
            pos = cursor.get(code_key, 1)
            treffer = next(
                (v for v in eintrag["vorkommen"] if v["position"] == pos),
                None
            )
            cursor[code_key] = pos + 1
            if treffer is None:
                fallback_count[code_key] = fallback_count.get(code_key, 0) + 1
                return eintrag["grundform"]
            return treffer["originalform"]
        # Schema v1: Grundform
        return eintrag["grundform"]
    return repl


def process_paragraph(paragraph, repl_callback) -> bool:
    """
    Bearbeitet einen Absatz. Gibt True zurueck, wenn der Absatz die
    Kennzeichnungszeile war und geloescht werden soll.
    """
    runs = paragraph.runs
    if not runs:
        return False
    text = "".join(run.text for run in runs)
    if MARKER_PATTERN.match(text.strip()):
        return True
    new_text = CODE_PATTERN.sub(repl_callback, text)
    if new_text != text:
        runs[0].text = new_text
        for run in runs[1:]:
            run.text = ""
    return False


def process_document(doc, schema_version: int, eintrag_by_code: dict[str, dict]):
    counter: dict[str, int] = {}
    missing: set[str] = set()
    cursor: dict[str, int] = {c: 1 for c in eintrag_by_code}
    fallback_count: dict[str, int] = {}
    repl_callback = make_reverse_replacer(
        schema_version, eintrag_by_code, cursor, counter, missing, fallback_count
    )

    to_remove: list = []

    def walk_container(container):
        for paragraph in container.paragraphs:
            if process_paragraph(paragraph, repl_callback):
                to_remove.append(paragraph)
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    walk_container(cell)

    walk_container(doc)
    for section in doc.sections:
        for hf in (section.header, section.footer,
                   section.first_page_header, section.first_page_footer,
                   section.even_page_header, section.even_page_footer):
            if hf is not None:
                walk_container(hf)

    for paragraph in to_remove:
        el = paragraph._element
        el.getparent().remove(el)

    return counter, missing, cursor, fallback_count


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Deterministische Rueckumwandlung eines PSEUDO-DOCX (Skill v1.1)."
    )
    parser.add_argument("--input", required=True, type=Path,
                        help="PSEUDO-DOCX (wird nicht veraendert)")
    parser.add_argument("--legend", required=True, type=Path,
                        help="JSON-Legende zum Alias")
    parser.add_argument("--outdir", required=True, type=Path,
                        help="Ausgabeverzeichnis")
    args = parser.parse_args()

    if not args.input.exists():
        sys.stderr.write(f"FEHLER: Eingabedatei nicht gefunden: {args.input}\n")
        return 2
    if not args.legend.exists():
        sys.stderr.write(f"FEHLER: Legende nicht gefunden: {args.legend}\n")
        return 2
    args.outdir.mkdir(parents=True, exist_ok=True)

    schema_version, eintrag_by_code = load_legend(args.legend)
    doc = Document(str(args.input))
    counter, missing, cursor, fallback_count = process_document(
        doc, schema_version, eintrag_by_code
    )

    stem = args.input.stem
    if stem.endswith("_PSEUDO"):
        klar_name = stem[: -len("_PSEUDO")] + "_KLAR.docx"
    else:
        klar_name = stem + "_KLAR.docx"
    klar_path = args.outdir / klar_name
    doc.save(str(klar_path))

    unused = [
        f"[{code}]"
        for code in eintrag_by_code
        if f"[{code}]" not in counter
    ]

    # Cursor-Konsistenzpruefung: Nach dem Lauf muss cursor[c] = len(vorkommen)+1 sein,
    # sonst wurden nicht alle Vorkommen abgerufen. Fuer v1 ist das Feld leer.
    cursor_konsistenz = {}
    if schema_version >= 2:
        for code, eintrag in eintrag_by_code.items():
            n = len(eintrag["vorkommen"])
            erwartet = n + 1
            ist = cursor.get(code, 1)
            if ist != erwartet:
                cursor_konsistenz[code] = {"erwartet": erwartet, "ist": ist}

    status = "OK"
    if missing:
        status = "WARNUNG: Codes ohne Legende gefunden"
    elif fallback_count:
        status = "WARNUNG: Positionsindex unvollstaendig, Fallback auf Grundform"
    elif cursor_konsistenz:
        status = "WARNUNG: Cursor-Konsistenz verletzt"
    elif schema_version < 2:
        status = "OK, aber Legendenschema v1: wortgleicher Roundtrip nicht garantiert"

    report = {
        "skill_version": "1.1",
        "legendenschema": schema_version,
        "eingabe": str(args.input),
        "legende": str(args.legend),
        "ausgabe_klar": str(klar_path),
        "ersetzte_codes": counter,
        "codes_ohne_legende": sorted(missing),
        "legendeneintraege_ohne_vorkommen": unused,
        "positionsindex_fallback_auf_grundform": fallback_count,
        "cursor_konsistenz_verletzungen": cursor_konsistenz,
        "status": status,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if missing or cursor_konsistenz:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
